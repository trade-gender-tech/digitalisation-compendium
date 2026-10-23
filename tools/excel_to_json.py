"""
Excel → data converter for the WTO Digitalization Compendium
(Empowering Women to Trade through Digitalization: Policy Insights Compendium).

Reads the compendium workbook and writes the two data files used by the page:

    downloads/digitalization-compendium.xlsx
        sheets "Questions", "WTO Members", "Other Stakeholders"
        → data/digitalization_data.json   the data, as loaded by index.html
        → data/digitalization_data.js     the same data as a script, used only when the page is opened
                                          from disk (browsers do not let a page opened from disk
                                          read a JSON file)

Cell values are copied verbatim. The only processing: multiple-choice cells (one option per
line) become lists of options, and dates become "YYYY-MM-DD hh:mm:ss" text.

Usage (from any folder):
    python3 tools/excel_to_json.py
    python3 tools/excel_to_json.py --input path/to/other.xlsx

Requirements:
    Python 3.8+ and openpyxl  (pip install -r tools/requirements.txt)

Run it every time the Excel file changes, then refresh index.html.
"""

import argparse
import json
import os
import re
import sys
from collections import Counter
from datetime import date, datetime

try:
    import openpyxl
except ImportError:
    sys.exit("ERROR: openpyxl is required.  Install it with:  pip install openpyxl")


PACKAGE_DIR    = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DEFAULT_INPUT  = os.path.join(PACKAGE_DIR, "downloads", "digitalization-compendium.xlsx")
DEFAULT_OUTDIR = os.path.join(PACKAGE_DIR, "data")
JS_VARIABLE    = "DIGITALIZATION_COMPENDIUM_DATA"   # index.html reads window.DIGITALIZATION_COMPENDIUM_DATA

TITLE = "Empowering Women to Trade through Digitalization: Policy Insights Compendium"

# Respondent sheets: sheet name → (group id, name of the added region/category column)
SHEETS = {
    "WTO Members":        ("members",      "Region"),
    "Other Stakeholders": ("stakeholders", "Category"),
}

# Column header "Q1 | …" (answer), "Q1.1 | …" or "S1.a | …" (explanation / details)
CODE_HEADER = re.compile(r"^(?P<code>[A-Z]\w*?)(?:\.(?P<part>1|[a-h]))? \| ")


def text(v):
    """Cell value → text, verbatim ('' for empty cells, dates as YYYY-MM-DD hh:mm:ss)."""
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


def read_questions(wb):
    ws = wb["Questions"]
    header = [text(c.value) for c in ws[1]]
    idx = {h: i for i, h in enumerate(header)}
    questions = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row[idx["Code"]]:
            continue
        prompts = text(row[idx["Explanation prompt(s)"]])
        questions.append({
            "code":     text(row[idx["Code"]]),
            "group":    "members" if text(row[idx["Respondents"]]) == "WTO Members" else "stakeholders",
            "theme":    text(row[idx["Theme"]]),
            "question": text(row[idx["Question"]]),
            "prompts":  [p for p in prompts.split("\n") if p.strip()],
            "kind":     "text" if text(row[idx["Answer type"]]).lower().startswith("free") else "options",
        })
    return questions


def read_respondents(ws, group, area_header, questions):
    """One sheet → respondents. Each respondent has:
         info     the response's own identification fields, verbatim (e.g. Time, Ministry/Institution:)
         answers  {code: {"options": [...] (option questions), "main": "..." (free-text questions),
                          "notes": ["...", ...] (explanations, in the order of the question's prompts)}}"""
    header = [text(c.value) for c in ws[1]]
    kinds = {q["code"]: q["kind"] for q in questions if q["group"] == group}
    respondents = []
    for values in ws.iter_rows(min_row=2, values_only=True):
        if not any(v not in (None, "") for v in values):
            continue
        r = {"id": "", "group": group, "name": "", "area": "", "info": [], "answers": {}}
        for h, v in zip(header, values):
            m = CODE_HEADER.match(h)
            if h == "Respondent":
                r["name"] = text(v)
            elif h == area_header:
                r["area"] = text(v)
            elif not m:
                r["info"].append({"label": h, "value": text(v)})
            elif m.group("code") not in kinds:
                print(f"WARNING: column '{h[:60]}' in '{ws.title}' has no entry in the Questions sheet; skipped")
            else:
                a = r["answers"].setdefault(m.group("code"), {"options": [], "main": "", "notes": []})
                if m.group("part"):
                    a["notes"].append(text(v))
                elif kinds[m.group("code")] == "text":
                    a["main"] = text(v)
                else:
                    a["options"] = [o.strip() for o in text(v).split("\n") if o.strip()]
        r["id"] = group[0] + "-" + re.sub(r"[^a-z0-9]+", "-", r["name"].lower()).strip("-")
        respondents.append(r)
    return respondents


def main():
    parser = argparse.ArgumentParser(description="Convert the Digitalization Compendium Excel file into data/digitalization_data.json and .js")
    parser.add_argument("--input",  "-i", default=DEFAULT_INPUT,  help="Excel file (default: downloads/digitalization-compendium.xlsx)")
    parser.add_argument("--outdir", "-o", default=DEFAULT_OUTDIR, help="output folder (default: data/)")
    args = parser.parse_args()
    if not os.path.isfile(args.input):
        sys.exit(f"ERROR: Input file not found: {args.input}")

    print(f"Reading:  {args.input}")
    wb = openpyxl.load_workbook(args.input, data_only=True)
    questions = read_questions(wb)
    respondents = []
    for sheet, (group, area_header) in SHEETS.items():
        respondents += read_respondents(wb[sheet], group, area_header, questions)

    themes = []
    for q in questions:
        if q["theme"] not in themes:
            themes.append(q["theme"])

    output = {
        "meta": {
            "title": TITLE,
            "generated_at": date.today().isoformat(),
            "groups": {"members": "WTO Members", "stakeholders": "Other stakeholders"},
            "themes": themes,
            "questions": questions,
        },
        "respondents": respondents,
    }
    json_str = json.dumps(output, ensure_ascii=False, indent=1)

    os.makedirs(args.outdir, exist_ok=True)
    json_path = os.path.join(args.outdir, "digitalization_data.json")
    js_path   = os.path.join(args.outdir, "digitalization_data.js")
    with open(json_path, "w", encoding="utf-8") as f:
        f.write(json_str + "\n")
    with open(js_path, "w", encoding="utf-8") as f:
        f.write("/* Digitalization Compendium data.\n"
                "   Generated from downloads/digitalization-compendium.xlsx by tools/excel_to_json.py: do not edit by hand. */\n")
        f.write(f"window.{JS_VARIABLE} = {json_str};\n")

    groups = Counter(r["group"] for r in respondents)
    print(f"Written:  {json_path}")
    print(f"Written:  {js_path}")
    print(f"Responses: {groups['members']} WTO Members, {groups['stakeholders']} other stakeholders; "
          f"{len(questions)} questions in {len(themes)} themes")


if __name__ == "__main__":
    main()
