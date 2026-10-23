"""
Questionnaire exports → compendium Excel, for the WTO Digitalization Compendium
(Empowering Women to Trade through Digitalization: Policy Insights Questionnaire).

Reads the two raw Formstack exports and writes the public compendium workbook:

    Members export       (one row per WTO Member response)
    Stakeholders export  (one row per other-stakeholder response)
        → downloads/digitalization-compendium.xlsx
              About                 title, source and notes
              WTO Members           one row per Member response
              Other Stakeholders    one row per other-stakeholder response
              Questions             every question: code, theme, wording, explanation prompt(s)

All answers, explanations and column headers are copied VERBATIM from the exports.
The only changes are:
  - personal data is left out: respondent names, positions, e-mails, phone numbers, browser,
    IP address, Formstack ID, location and links to uploaded files (Formstack admin links).
    Where a respondent wrote personal contact details inside an answer, they are replaced by
    "[contact details removed]".
  - two columns are added at the front of each sheet: "Respondent" (the Member's standard WTO
    name, or the organization's name) and "Region" (Members) / "Category" (other stakeholders).
  - each question column header is prefixed with a code, e.g. "Q1 | 1. Do you explicitly…",
    and each explanation column with the question code + ".1" (Members) or ".a"/".b"
    (other stakeholders), so the converter can match answers to questions.

Usage:
    python3 tools/questionnaire_to_excel.py --members MEMBERS.xlsx --stakeholders STAKEHOLDERS.xlsx

Then run tools/excel_to_json.py to rebuild the data files used by the page.
When a new Member responds, add it to MEMBERS (name and region) below if the script asks for it.
"""

import argparse
import os
import re
import sys
from datetime import date, datetime

try:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
except ImportError:
    sys.exit("ERROR: openpyxl is required.  Install it with:  pip install openpyxl")


PACKAGE_DIR    = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DEFAULT_OUTPUT = os.path.join(PACKAGE_DIR, "downloads", "digitalization-compendium.xlsx")

TITLE = "Empowering Women to Trade through Digitalization: Policy Insights Compendium"
REDACTED = "[contact details removed]"

# ---------------------------------------------------------------------------
# Respondents: name as submitted → (standard name, region or category)
# ---------------------------------------------------------------------------
MEMBERS = {
    "Nigeria":                         ("Nigeria",            "Africa"),
    "Government of Mauritius":         ("Mauritius",          "Africa"),
    "The Philippines":                 ("Philippines",        "Asia"),
    "China":                           ("China",              "Asia"),
    "WTO Member/Bulgaria":             ("Bulgaria",           "Europe"),
    "Ukraine":                         ("Ukraine",            "Europe"),
    "United Kingdom":                  ("United Kingdom",     "Europe"),
    "SLOVENIA, The Republic of":       ("Slovenia",           "Europe"),
    "ICEX SPAIN TRADE AND INVESTMENT": ("Spain",              "Europe"),
    "Russian Federation":              ("Russian Federation", "Europe"),
    "Brasil":                          ("Brazil",             "Latin America and the Caribbean"),
    "Chile":                           ("Chile",              "Latin America and the Caribbean"),
    "Colombia":                        ("Colombia",           "Latin America and the Caribbean"),
    "Dominican Republic":              ("Dominican Republic", "Latin America and the Caribbean"),
    "Ecuador":                         ("Ecuador",            "Latin America and the Caribbean"),
    "México":                          ("Mexico",             "Latin America and the Caribbean"),
    "Peru":                            ("Peru",               "Latin America and the Caribbean"),
    "Australia":                       ("Australia",          "Oceania"),
}
# Responses whose Member field does not name the Member ("WTO MEMBER"): identified by the ministry
MEMBER_BY_INSTITUTION = {
    "Ministry of Economic Development of the Russian Federation": "Russian Federation",
}
STAKEHOLDERS = {
    "ILO":                               ("International Labour Organization (ILO)", "International organization"),
    "UN Trade and Development (UNCTAD)": ("UN Trade and Development (UNCTAD)", "International organization"),
    "Universal postal union":            ("Universal Postal Union (UPU)", "International organization"),
    "Institute of International Studies, University of Chile":
                                         ("Institute of International Studies, University of Chile", "Academia"),
}

# Raw columns that are personal or technical data and are left out (matched on the header)
PERSONAL_COLUMNS = {"name", "position:", "email:", "phone number/whatsapp:", "upload supporting documents (optional)",
                    "focal point full name:", "email", "phone/whatsapp",
                    "browser", "ip address", "unique id", "location"}

# ---------------------------------------------------------------------------
# Members questionnaire: code → theme and a phrase from the question's column header
# (columns are found by phrase, so the script keeps working if Formstack reorders them).
# The explanation column of question N is the one whose header starts with "N.1.".
# Themes group the questions for navigation; they follow the order of the questionnaire.
# ---------------------------------------------------------------------------
THEMES = [
    "Strategies and priorities",
    "Initiatives, instruments and market access",
    "Skills and capacity building",
    "Coordination and actors",
    "Data, monitoring and evaluation",
    "Trade facilitation, innovation and finance",
    "Additional information",
]
STAKEHOLDER_THEME = "Other stakeholders questionnaire"

MEMBER_QUESTIONS = [
    # code,    theme, header phrase
    ("Q1",     0, "explicitly include women in your national digital"),
    ("Q2",     0, "which dimensions of digitalization"),
    ("Q3",     0, "greatest opportunities for women"),
    ("Q4",     0, "greatest challenges for women"),
    ("Q5",     0, "which policy objectives do you currently prioritize"),
    ("Q6",     0, "what are your priority actions"),
    ("Q7",     1, "do you implement initiatives or programmes specifically supporting women"),
    ("Q8",     1, "which areas of intervention"),
    ("Q9",     1, "trade policies that support women with digitalization"),
    ("Q10",    1, "which instruments do you use"),
    ("Q11",    1, "to what extent do policies or programmes help"),
    ("Q12",    1, "which forms of market access support"),
    ("Q13",    2, "do you offer training on digital entrepreneurship"),
    ("Q14",    2, "which skills are currently covered"),
    ("Q15",    2, "train-the-trainer"),
    ("Q15bis", 2, "have you observed a measurable impact"),
    ("Q16",    2, "capacity-building programmes in partnership"),
    ("Q17",    3, "how well are policies related to women"),
    ("Q18",    3, "governance mechanisms in place"),
    ("Q19",    3, "which actors are involved"),
    ("Q20",    4, "do you collect sex-disaggregated data"),
    ("Q21",    4, "which types of sex-disaggregated data"),
    ("Q22",    4, "which institutions or entities are involved"),
    ("Q23",    4, "data on barriers affecting women"),
    ("Q24",    4, "qualitative and/or quantitative data"),
    ("Q25",    4, "monitored and evaluated"),
    ("Q26",    5, "link digital infrastructure development"),
    ("Q27",    5, "innovation-driven trade initiatives"),
    ("Q28",    5, "accessing investment, start-up funding"),
    ("Q29",    6, "please provide any additional comments"),
]
TEXT_QUESTIONS = {"Q29"}   # free-text answers (all other Member questions are answer options)


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def norm(h):
    return re.sub(r"\s+", " ", str(h or "")).strip().lower()


def header_text(h):
    """Raw header with tabs and repeated spaces shown as single spaces (wording unchanged)."""
    return re.sub(r"[ \t]+", " ", str(h or "")).strip()


def redact(value, names):
    """Replace personal contact details written inside an answer; everything else is unchanged."""
    if not isinstance(value, str):
        return value
    s = value
    for name in names:
        if name and name in s:
            s = s.replace(name, REDACTED)
    s = re.sub(r"[\w.+-]+@[\w-]+(\.[\w-]+)+", REDACTED, s)         # e-mail addresses
    s = re.sub(r"\+\d[\d\s().-]{7,}\d", REDACTED, s)               # international phone numbers
    s = re.sub(re.escape(REDACTED) + r"(\s*,\s*" + re.escape(REDACTED) + r")+", REDACTED, s)  # one marker per run
    return s


def is_personal(header):
    return norm(header) in PERSONAL_COLUMNS


# ---------------------------------------------------------------------------
# Readers
# ---------------------------------------------------------------------------

def read_members(path):
    ws = openpyxl.load_workbook(path, data_only=True).active
    raw = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    heads = [norm(h) for h in raw]

    def find(phrase):
        hits = [i for i, h in enumerate(heads) if phrase.lower() in h]
        if not hits:
            sys.exit(f"ERROR: Members export: no column header contains '{phrase}'.")
        return hits[0]

    member_col, inst_col, name_col = find("wto member/observer"), find("ministry/institution"), heads.index("name")
    info_cols = [i for i, h in enumerate(heads) if h in ("time", "wto member/observer governement:", "ministry/institution:")]

    questions, columns = [], []   # columns: (raw index, output header)
    for c in info_cols:
        columns.append((c, header_text(raw[c])))
    for code, theme, phrase in MEMBER_QUESTIONS:
        a = find(phrase)
        number = code[1:].replace("bis", ".bis")
        note = next((i for i, h in enumerate(heads) if h.startswith(number + ".1.")), None) if code not in TEXT_QUESTIONS else None
        questions.append({"code": code, "theme": THEMES[theme], "question": header_text(raw[a]),
                          "prompts": [header_text(raw[note])] if note is not None else [],
                          "kind": "Free text" if code in TEXT_QUESTIONS else "Answer options"})
        columns.append((a, f"{code} | {header_text(raw[a])}"))
        if note is not None:
            columns.append((note, f"{code}.1 | {header_text(raw[note])}"))

    used = {c for c, _ in columns}
    left_out = [header_text(raw[i]) for i in range(len(raw)) if i not in used]
    unexpected = [h for h in left_out if not is_personal(h)]
    if unexpected:
        sys.exit("ERROR: Members export has columns this script does not know: " + "; ".join(unexpected))

    rows, unknown = [], []
    for r in range(2, ws.max_row + 1):
        vals = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
        if not any(v not in (None, "") for v in vals):
            continue
        raw_member, raw_inst = str(vals[member_col] or "").strip(), str(vals[inst_col] or "").strip()
        if raw_inst in MEMBER_BY_INSTITUTION:
            name, region = MEMBERS[MEMBER_BY_INSTITUTION[raw_inst]]
        elif raw_member in MEMBERS:
            name, region = MEMBERS[raw_member]
        else:
            unknown.append(raw_member)
            continue
        names = [str(vals[name_col] or "").strip()]
        rows.append([name, region] + [redact(vals[c], names) for c, _ in columns])
    if unknown:
        sys.exit("ERROR: add these Members to MEMBERS in this script: " + ", ".join(repr(u) for u in unknown))
    rows.sort(key=lambda x: x[0])
    return questions, ["Respondent", "Region"] + [h for _, h in columns], rows


def read_stakeholders(path):
    """Questions start with a number ("1. Have governments…"); the columns after each one,
    up to the next question, hold its details ("Please provide details…")."""
    ws = openpyxl.load_workbook(path, data_only=True).active
    raw = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    heads = [norm(h) for h in raw]
    time_col = heads.index("time")
    org_col = next(i for i, h in enumerate(heads) if h.startswith("organisation/institution"))
    name_col = next(i for i, h in enumerate(heads) if h.startswith("focal point full name"))

    columns = [(time_col, header_text(raw[time_col])), (org_col, header_text(raw[org_col]))]
    questions, current, n = [], None, 0
    for i, h in enumerate(heads):
        if i in (time_col, org_col) or is_personal(raw[i]):
            continue
        if re.match(r"^\d+\.", h):
            n += 1
            current = {"code": f"S{n}", "theme": STAKEHOLDER_THEME, "question": header_text(raw[i]), "prompts": []}
            questions.append(current)
            columns.append((i, f"S{n} | {header_text(raw[i])}"))
        elif current is not None:
            letter = "abcdefgh"[len(current["prompts"])]
            current["prompts"].append(header_text(raw[i]))
            columns.append((i, f"{current['code']}.{letter} | {header_text(raw[i])}"))
        else:
            sys.exit(f"ERROR: Stakeholders export: unexpected column '{header_text(raw[i])}'.")

    # A question whose own column holds short fixed answers (Yes / No / Under development…) is an
    # answer-options question; one with long written answers is free text.
    for q in questions:
        c = next(c for c, h in columns if h.startswith(q["code"] + " | "))
        values = [ws.cell(r, c + 1).value for r in range(2, ws.max_row + 1)]
        q["kind"] = "Answer options" if all(v is None or len(str(v)) <= 40 for v in values) else "Free text"

    rows = []
    for r in range(2, ws.max_row + 1):
        vals = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
        raw_org = str(vals[org_col] or "").strip()
        if not raw_org:
            continue
        if raw_org not in STAKEHOLDERS:
            sys.exit(f"ERROR: add {raw_org!r} to STAKEHOLDERS in this script.")
        name, category = STAKEHOLDERS[raw_org]
        names = [str(vals[name_col] or "").strip()]
        rows.append([name, category] + [redact(vals[c], names) for c, _ in columns])
    rows.sort(key=lambda x: x[0])
    return questions, ["Respondent", "Category"] + [h for _, h in columns], rows


# ---------------------------------------------------------------------------
# Writer
# ---------------------------------------------------------------------------

HEAD_FONT = Font(bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="004080")
WRAP_TOP  = Alignment(wrap_text=True, vertical="top")


def write_sheet(ws, header, rows):
    ws.append(header)
    for row in rows:
        ws.append(row)
    for c, h in enumerate(header, start=1):
        cell = ws.cell(1, c)
        cell.font, cell.fill, cell.alignment = HEAD_FONT, HEAD_FILL, WRAP_TOP
        explanation = bool(re.match(r"^\w+\.(1|[a-h]) \| ", h))
        ws.column_dimensions[cell.column_letter].width = 22 if c <= 4 else (55 if explanation else 34)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = WRAP_TOP
            if isinstance(cell.value, datetime):
                cell.number_format = "yyyy-mm-dd hh:mm:ss"
    ws.row_dimensions[1].height = 90
    ws.freeze_panes = ws.cell(2, 2)
    ws.auto_filter.ref = ws.dimensions


def write_workbook(path, m_questions, m_header, members, s_questions, s_header, stakeholders, period):
    wb = openpyxl.Workbook()
    about = wb.active
    about.title = "About"
    lines = [
        (TITLE, Font(bold=True, size=14, color="004080")),
        ("WTO Trade and Gender Office", Font(italic=True, color="555555")),
        ("", None),
        (f"Responses: {len(members)} WTO Members and {len(stakeholders)} other stakeholders, received {period}.", None),
        ("'WTO Members' and 'Other Stakeholders' hold one row per response. Answers, explanations and question "
         "wording are reproduced verbatim, as submitted, in the language of the respondent.", None),
        ("Multiple-choice answers are listed one per line. Each question column starts with a code (e.g. 'Q1 |'); "
         "its explanation column has the same code followed by '.1' (Members) or '.a' / '.b' (other stakeholders). "
         "'Questions' lists every question with its theme.", None),
        ("The 'Respondent' and 'Region' / 'Category' columns were added for reference. Personal data (names, "
         "positions and contact details of respondents) has been removed; contact details written inside an "
         f"answer are shown as {REDACTED}.", None),
        (f"Compiled on {date.today().isoformat()}.", Font(color="777777")),
    ]
    for text, font in lines:
        about.append([text])
        if font:
            about.cell(about.max_row, 1).font = font
        about.cell(about.max_row, 1).alignment = Alignment(wrap_text=True, vertical="top")
    about.column_dimensions["A"].width = 120

    write_sheet(wb.create_sheet("WTO Members"), m_header, members)
    write_sheet(wb.create_sheet("Other Stakeholders"), s_header, stakeholders)

    qs = wb.create_sheet("Questions")
    qs.append(["Code", "Respondents", "Theme", "Question", "Explanation prompt(s)", "Answer type"])
    for group, questions in (("WTO Members", m_questions), ("Other stakeholders", s_questions)):
        for q in questions:
            qs.append([q["code"], group, q["theme"], q["question"], "\n".join(q["prompts"]), q["kind"]])
    for c, w in enumerate([9, 18, 36, 90, 60, 16], start=1):
        qs.column_dimensions[qs.cell(1, c).column_letter].width = w
        qs.cell(1, c).font, qs.cell(1, c).fill = HEAD_FONT, HEAD_FILL
    for row in qs.iter_rows():
        for cell in row:
            cell.alignment = WRAP_TOP
    qs.freeze_panes = "A2"

    os.makedirs(os.path.dirname(path), exist_ok=True)
    wb.save(path)


def main():
    parser = argparse.ArgumentParser(description="Build downloads/digitalization-compendium.xlsx from the questionnaire exports")
    parser.add_argument("--members", required=True, help="Formstack export of the WTO Members questionnaire (.xlsx)")
    parser.add_argument("--stakeholders", required=True, help="Formstack export of the other-stakeholders questionnaire (.xlsx)")
    parser.add_argument("--output", "-o", default=DEFAULT_OUTPUT, help="output workbook (default: downloads/digitalization-compendium.xlsx)")
    parser.add_argument("--period", default="between June and October 2026", help="text describing when responses were received")
    args = parser.parse_args()

    m_questions, m_header, members = read_members(args.members)
    s_questions, s_header, stakeholders = read_stakeholders(args.stakeholders)
    write_workbook(args.output, m_questions, m_header, members, s_questions, s_header, stakeholders, args.period)
    print(f"Written:  {args.output}")
    print(f"WTO Members:        {len(members)} responses, {len(m_questions)} questions")
    print(f"Other stakeholders: {len(stakeholders)} responses, {len(s_questions)} questions")


if __name__ == "__main__":
    main()
